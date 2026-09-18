import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def safe_name(value):
    """Convert text into a filesystem-safe lowercase name."""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def clean_sequence(sequence):
    """Remove whitespace and force uppercase."""
    return re.sub(r"\s+", "", str(sequence)).upper()


def normalize_serial(value):
    """Normalize Excel serial values like 10000393.0 to 10000393."""
    if value is None:
        return None

    try:
        return int(float(value))
    except (ValueError, TypeError):
        return str(value).strip()


def main():
    parser = argparse.ArgumentParser(
        description="Generate an AlphaFold 3 input JSON from the prepared UTexas dataset."
    )

    parser.add_argument(
        "--excel",
        required=True,
        help="Path to the prepared Excel spreadsheet",
    )

    parser.add_argument(
        "--serial",
        required=True,
        help="Serial Number of the aptamer entry",
    )

    parser.add_argument(
        "--output-dir",
        default="af3_inputs",
        help="Directory where AF3 JSON files will be written",
    )

    args = parser.parse_args()

    workbook = load_workbook(args.excel, data_only=True)

    if "Prepared" not in workbook.sheetnames:
        raise ValueError(
            f"'Prepared' sheet not found. Sheets present: {workbook.sheetnames}"
        )

    sheet = workbook["Prepared"]

    # Build mapping: stripped column header -> column number
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }

    required_columns = [
        "Serial Number",
        "Target",
        "Name of Aptamer",
        "Clean Sequence",
        "AF3 Molecule Type",
        "Protein Sequence(s)",
        "Target Mapping Status",
    ]

    missing_columns = [
        column for column in required_columns
        if column not in headers
    ]

    if missing_columns:
        raise ValueError(
            f"Missing required spreadsheet columns: {missing_columns}"
        )

    requested_serial = normalize_serial(args.serial)

    selected_row = None

    # Find the requested entry by Serial Number
    for row in sheet.iter_rows(min_row=2):
        serial_value = row[headers["Serial Number"] - 1].value
        normalized_value = normalize_serial(serial_value)

        if normalized_value == requested_serial:
            selected_row = row
            break

    if selected_row is None:
        raise ValueError(
            f"Serial Number {args.serial} was not found in the Prepared sheet."
        )

    def get(column_name):
        return selected_row[headers[column_name] - 1].value

    serial = normalize_serial(get("Serial Number"))
    target = get("Target")
    aptamer_name = get("Name of Aptamer")
    aptamer_sequence = get("Clean Sequence")
    molecule_type = str(get("AF3 Molecule Type")).strip().lower()
    protein_sequences_raw = get("Protein Sequence(s)")
    mapping_status = str(get("Target Mapping Status")).strip()

    # Validation
    if mapping_status != "Mapped":
        raise ValueError(
            f"Target Mapping Status is '{mapping_status}', not 'Mapped'."
        )

    if molecule_type not in {"dna", "rna"}:
        raise ValueError(
            f"Unsupported AF3 molecule type: '{molecule_type}'"
        )

    if not aptamer_sequence:
        raise ValueError("Clean Sequence is empty.")

    if not protein_sequences_raw:
        raise ValueError("Protein Sequence(s) is empty.")

    # Split multiple protein chains using |
    protein_sequences = [
        clean_sequence(sequence)
        for sequence in str(protein_sequences_raw).split("|")
        if sequence.strip()
    ]

    if not protein_sequences:
        raise ValueError("No valid protein sequences were found.")

    sequences = []

    # Assign A, B, C... to protein chains
    for index, protein_sequence in enumerate(protein_sequences):
        if index >= 26:
            raise ValueError(
                "More than 26 protein chains are not currently supported by this script."
            )

        chain_id = chr(ord("A") + index)

        sequences.append({
            "protein": {
                "id": chain_id,
                "sequence": protein_sequence
            }
        })

    # Aptamer gets the next chain ID
    aptamer_chain_index = len(protein_sequences)

    if aptamer_chain_index >= 26:
        raise ValueError(
            "No available single-letter chain ID for the aptamer."
        )

    aptamer_chain_id = chr(ord("A") + aptamer_chain_index)

    sequences.append({
        molecule_type: {
            "id": aptamer_chain_id,
            "sequence": clean_sequence(aptamer_sequence)
        }
    })

    # Standard naming convention:
    # <serial>_<target>_<aptamer>
    job_name = (
        f"{serial}_"
        f"{safe_name(target)}_"
        f"{safe_name(aptamer_name)}"
    )

    af3_input = {
        "name": job_name,
        "modelSeeds": [1],
        "sequences": sequences,
        "dialect": "alphafold3",
        "version": 4
    }

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    output_path = output_dir / f"{job_name}.json"

    with open(output_path, "w") as file:
        json.dump(af3_input, file, indent=2)

    print(f"Created: {output_path}")
    print(f"Job name: {job_name}")
    print(f"Protein chains: {len(protein_sequences)}")
    print(f"Aptamer type: {molecule_type}")
    print(f"Aptamer chain: {aptamer_chain_id}")


if __name__ == "__main__":
    main()
