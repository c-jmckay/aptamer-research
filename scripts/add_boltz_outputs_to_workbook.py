import csv
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path(
    "data/UTexas_Aptamer_Database_with_AF3_and_Boltz_outputs.xlsx"
)

RESULTS_PATH = Path(
    "results/boltz_results_combined.csv"
)

OUTPUT_PATH = Path(
    "data/UTexas_Aptamer_Database_expanded_results.xlsx"
)

SOURCE_SHEET = "Prepared"
OUTPUT_SHEET = "Boltz Outputs"


def normalize_serial(value):
    if value is None:
        return None

    try:
        return str(int(float(value)))
    except (ValueError, TypeError):
        return str(value).strip()


def to_float(value):
    if value in (None, ""):
        return None

    try:
        return float(value)
    except (ValueError, TypeError):
        return None


# ------------------------------------------------------------
# Load workbook
# ------------------------------------------------------------

wb = load_workbook(WORKBOOK_PATH)

source = wb[SOURCE_SHEET]

headers = {
    str(cell.value).strip(): idx
    for idx, cell in enumerate(source[1], start=1)
    if cell.value is not None
}

required_columns = [
    "Serial Number",
    "Target",
    "Name of Aptamer",
    "Kd (nM)",
]

for column in required_columns:
    if column not in headers:
        raise ValueError(
            f"Missing required column in Prepared sheet: {column}"
        )


# ------------------------------------------------------------
# Build UTexas metadata lookup
# ------------------------------------------------------------

metadata = {}

for row in source.iter_rows(min_row=2):

    serial = normalize_serial(
        row[headers["Serial Number"] - 1].value
    )

    if serial is None:
        continue

    metadata[serial] = {
        "Target": row[headers["Target"] - 1].value,
        "Name of Aptamer":
            row[headers["Name of Aptamer"] - 1].value,
        "Kd (nM)":
            row[headers["Kd (nM)"] - 1].value,
    }


# ------------------------------------------------------------
# Read Boltz results
# ------------------------------------------------------------

with RESULTS_PATH.open(newline="") as f:
    reader = csv.DictReader(f)
    boltz_rows = list(reader)


# Sort complete predictions from lowest affinity value to highest.
# Lower Boltz affinity_pred_value = stronger predicted binding.

boltz_rows.sort(
    key=lambda row: (
        to_float(row["affinity_pred_value"]) is None,
        (
            to_float(row["affinity_pred_value"])
            if to_float(row["affinity_pred_value"]) is not None
            else float("inf")
        ),
    )
)


# ------------------------------------------------------------
# Replace old Boltz sheet if present
# ------------------------------------------------------------

if OUTPUT_SHEET in wb.sheetnames:
    del wb[OUTPUT_SHEET]

ws = wb.create_sheet(OUTPUT_SHEET)


# ------------------------------------------------------------
# Headers
# ------------------------------------------------------------

output_headers = [
    "Predicted Affinity Rank",
    "Serial Number",
    "Target",
    "Name of Aptamer",
    "Kd (nM)",
    "Boltz Job Name",
    "Status",
    "Affinity Pred Value",
    "Affinity Pred Value 1",
    "Affinity Pred Value 2",
    "Affinity Probability Binary",
    "Affinity Probability Binary 1",
    "Affinity Probability Binary 2",
]

ws.append(output_headers)


# ------------------------------------------------------------
# Write data
# ------------------------------------------------------------

rank = 0

for result in boltz_rows:

    serial = normalize_serial(result["serial_number"])
    info = metadata.get(serial, {})

    affinity_value = to_float(
        result["affinity_pred_value"]
    )

    if result["status"] == "complete" and affinity_value is not None:
        rank += 1
        prediction_rank = rank
    else:
        prediction_rank = ""

    ws.append([
        prediction_rank,
        serial,
        info.get("Target", ""),
        info.get("Name of Aptamer", ""),
        info.get("Kd (nM)", ""),
        result["job_name"],
        result["status"],
        affinity_value,
        to_float(result["affinity_pred_value1"]),
        to_float(result["affinity_pred_value2"]),
        to_float(result["affinity_probability_binary"]),
        to_float(result["affinity_probability_binary1"]),
        to_float(result["affinity_probability_binary2"]),
    ])


# ------------------------------------------------------------
# Formatting
# ------------------------------------------------------------

header_fill = PatternFill(
    "solid",
    fgColor="7030A0"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions


# Numeric formats

for row_num in range(2, ws.max_row + 1):

    # Kd
    ws.cell(row_num, 5).number_format = "0.########"

    # Affinity prediction values
    for col in range(8, 11):
        ws.cell(row_num, col).number_format = "0.0000"

    # Binary probabilities
    for col in range(11, 14):
        ws.cell(row_num, col).number_format = "0.0000"


# Column widths

widths = {
    1: 22,
    2: 14,
    3: 40,
    4: 38,
    5: 14,
    6: 58,
    7: 14,
    8: 20,
    9: 20,
    10: 20,
    11: 28,
    12: 28,
    13: 28,
}

for col_idx, width in widths.items():
    ws.column_dimensions[
        get_column_letter(col_idx)
    ].width = width


for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


# ------------------------------------------------------------
# Save
# ------------------------------------------------------------

wb.save(OUTPUT_PATH)

print(
    f"Wrote {len(boltz_rows)} Boltz runs "
    f"to '{OUTPUT_SHEET}'"
)

print(
    f"Ranked {rank} complete affinity predictions "
    f"from low to high."
)

print(
    f"Saved workbook to: {OUTPUT_PATH}"
)
