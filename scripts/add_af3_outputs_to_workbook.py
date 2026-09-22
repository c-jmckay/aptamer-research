import csv
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path("data/UTexas_Aptamer_Database_expanded_results.xlsx")
RESULTS_PATH = Path("results/af3_results_combined.csv")
OUTPUT_PATH = Path("data/UTexas_Aptamer_Database_expanded_results_final.xlsx")

SOURCE_SHEET = "Prepared"
OUTPUT_SHEET = "AF3 Outputs"


def normalize_serial(value):
    if value is None:
        return None

    try:
        return str(int(float(value)))
    except (ValueError, TypeError):
        return str(value).strip()


# ------------------------------------------------------------
# Load workbook and source metadata
# ------------------------------------------------------------

wb = load_workbook(WORKBOOK_PATH)
source = wb[SOURCE_SHEET]

headers = {
    str(cell.value).strip(): idx
    for idx, cell in enumerate(source[1], start=1)
    if cell.value is not None
}

required = [
    "Serial Number",
    "Target",
    "Name of Aptamer",
    "Kd (nM)",
]

for column in required:
    if column not in headers:
        raise ValueError(f"Missing required column in Prepared sheet: {column}")


# Build lookup by serial number.
metadata = {}

for row in source.iter_rows(min_row=2):
    serial = normalize_serial(
        row[headers["Serial Number"] - 1].value
    )

    if serial is None:
        continue

    metadata[serial] = {
        "Target": row[headers["Target"] - 1].value,
        "Name of Aptamer": row[headers["Name of Aptamer"] - 1].value,
        "Kd (nM)": row[headers["Kd (nM)"] - 1].value,
    }


# ------------------------------------------------------------
# Read AF3 results CSV
# ------------------------------------------------------------

with RESULTS_PATH.open(newline="") as f:
    reader = csv.DictReader(f)
    af3_rows = list(reader)


# ------------------------------------------------------------
# Replace AF3 Outputs sheet if it already exists
# ------------------------------------------------------------

if OUTPUT_SHEET in wb.sheetnames:
    del wb[OUTPUT_SHEET]

ws = wb.create_sheet(OUTPUT_SHEET)


# ------------------------------------------------------------
# Output columns
# ------------------------------------------------------------

output_headers = [
    "Serial Number",
    "Target",
    "Name of Aptamer",
    "Kd (nM)",
    "AF3 Job Name",
    "Status",
    "Aptamer Type",
    "Protein Chains",
    "Aptamer Chain",
    "Sample Count",
    "pTM",
    "ipTM",
    "Ranking Score",
    "Fraction Disordered",
    "Has Clash",
    "Aptamer Chain pTM",
    "Aptamer Chain ipTM",
    "Best Aptamer-Protein ipTM",
    "Min Aptamer-Protein PAE",
    "Interface ipTM by Protein Chain",
    "Protein → Aptamer PAE by Chain",
    "Aptamer → Protein PAE by Chain",
]

ws.append(output_headers)


# ------------------------------------------------------------
# Write one row per AF3 run
# ------------------------------------------------------------

for result in af3_rows:
    serial = normalize_serial(result["serial_number"])
    info = metadata.get(serial, {})

    ws.append([
        serial,
        info.get("Target", ""),
        info.get("Name of Aptamer", ""),
        info.get("Kd (nM)", ""),
        result["job_name"],
        result["status"],
        result["aptamer_type"],
        result["protein_chains"],
        result["aptamer_chain"],
        result["sample_count"],
        result["ptm"],
        result["iptm"],
        result["ranking_score"],
        result["fraction_disordered"],
        result["has_clash"],
        result["aptamer_chain_ptm"],
        result["aptamer_chain_iptm"],
        result["best_aptamer_protein_iptm"],
        result["min_aptamer_protein_pae"],
        result["interface_iptm_by_chain"],
        result["protein_to_aptamer_pae_by_chain"],
        result["aptamer_to_protein_pae_by_chain"],
    ])


# ------------------------------------------------------------
# Formatting
# ------------------------------------------------------------

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# Numeric formatting.
for row in range(2, ws.max_row + 1):
    ws.cell(row, 4).number_format = "0.########"      # Kd
    for col in range(11, 20):
        ws.cell(row, col).number_format = "0.000"

# Sensible widths.
widths = {
    1: 14,
    2: 35,
    3: 35,
    4: 14,
    5: 55,
    6: 14,
    7: 14,
    8: 18,
    9: 14,
    10: 14,
    11: 12,
    12: 12,
    13: 14,
    14: 18,
    15: 12,
    16: 18,
    17: 18,
    18: 24,
    19: 24,
    20: 40,
    21: 40,
    22: 40,
}

for col_idx, width in widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


# ------------------------------------------------------------
# Save new workbook
# ------------------------------------------------------------

wb.save(OUTPUT_PATH)

print(f"Wrote {len(af3_rows)} AF3 runs to '{OUTPUT_SHEET}'")
print(f"Saved workbook to: {OUTPUT_PATH}")
